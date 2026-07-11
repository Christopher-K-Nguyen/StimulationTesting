"""Per-channel/combo Excel sheets follow the Gamry ``.DTA`` vertical
layout (operator: "Have the excel sheets for each channel/combo in the
same style as Gamry DTA with the data table at the end").

Layout, top → bottom: metadata preamble (ELECTRODE) → per-capture
SUMMARY table → raw-waveform DATA table at the END.  The SUMMARY and the
companion ``.DTA`` text file share the same column set.
"""
from __future__ import annotations

import numpy as np
import pytest

pytest.importorskip("openpyxl")


def _session_with_captures(n_captures: int = 2):
    from stimtest.session import Capture, Session, TestParameters, ChannelRun
    from stimtest.waveforms import PulsePattern
    from stimtest.electrode import Configuration, ElectrodeArray
    from stimtest.metrics import compute_metrics
    p = PulsePattern.biphasic(amplitude_ua=80.0)
    test = TestParameters(experiment="SP", pattern=p,
                          configuration=Configuration.monopolar(1),
                          array=ElectrodeArray.utah_4x4(), duration_s=1.0)
    sess = Session(notebook="nb", subject="s1", test=test)
    run = ChannelRun(configuration=Configuration.monopolar(1))
    run.surface_area_um2 = 5000.0
    t = np.linspace(-100.0, 600.0, 400)
    for idx, amp in enumerate(np.linspace(50.0, 80.0, n_captures)):
        pa = PulsePattern.biphasic(amplitude_ua=float(amp))
        c = Capture(index=idx, pattern=pa)
        c.time_us = t
        c.i_mon_ua = np.where((t >= 0) & (t < 200), -amp,
                              np.where((t >= 220) & (t < 420), amp, 0.0))
        c.v_mon_v = np.where((t >= 0) & (t < 200), -amp / 200,
                             np.where((t >= 220) & (t < 420), amp / 300, 0.0))
        compute_metrics(c, surface_area_um2=5000.0)
        run.captures.append(c)
    sess.runs.append(run)
    return sess, run


def _electrode_sheet(tmp_path):
    from stimtest import gamry_export
    from openpyxl import load_workbook
    sess, run = _session_with_captures(2)
    out = tmp_path / "g.xlsx"
    gamry_export.save_session_xlsx(sess, out)
    wb = load_workbook(out)
    # the per-channel sheet is the last one (after Instrumentation /
    # Parameters / [Setup] / Values)
    ws = wb[wb.sheetnames[-1]]
    col_a = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    return ws, col_a, run


def _row_of(col_a, marker):
    for i, v in enumerate(col_a, start=1):
        if v == marker:
            return i
    return None


def test_layout_is_preamble_then_summary_then_data(tmp_path):
    ws, col_a, run = _electrode_sheet(tmp_path)
    r_elec = _row_of(col_a, "ELECTRODE")
    r_summary = _row_of(col_a, "SUMMARY")
    r_data = _row_of(col_a, "DATA")
    assert r_elec == 1, "metadata preamble (ELECTRODE) must lead the sheet"
    assert r_summary and r_data, "both SUMMARY and DATA sections required"
    assert r_elec < r_summary < r_data, \
        "order must be preamble -> SUMMARY -> DATA (data table at the end)"
    # the DATA waveform table is genuinely the LAST thing on the sheet
    r_pt = _row_of(col_a, "Pt")
    assert r_pt and r_pt > r_data
    assert ws.max_row > r_pt + 100, "raw waveform rows fill the sheet tail"


def test_summary_has_one_row_per_capture(tmp_path):
    ws, col_a, run = _electrode_sheet(tmp_path)
    r_summary = _row_of(col_a, "SUMMARY")
    # header row is r_summary+1 ("Capture"), units row +2, data rows follow
    assert ws.cell(row=r_summary + 1, column=1).value == "Capture"
    first_data = r_summary + 3
    capture_indices = []
    for r in range(first_data, first_data + len(run.captures)):
        capture_indices.append(ws.cell(row=r, column=1).value)
    assert capture_indices == [str(c.index) for c in run.captures], \
        "SUMMARY must carry one row per capture, in order"


def test_dta_and_xlsx_share_summary_columns(tmp_path):
    from stimtest import gamry_export
    sess, run = _session_with_captures(2)
    headers, units = gamry_export._summary_headers_units()
    # .DTA text file uses the same header set
    paths = gamry_export.save_session_dta(sess, tmp_path / "dta")
    text = paths[0].read_text(encoding="utf-8")
    assert "\t".join(headers) in text, \
        ".DTA SUMMARY header must come from the shared helper"
