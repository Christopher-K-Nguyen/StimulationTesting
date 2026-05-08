#!/usr/bin/env python
"""Print the structure of an exported .xlsx workbook for inspection."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: inspect_xlsx.py <path.xlsx> [max_rows]")
        return 1
    path = Path(sys.argv[1])
    max_rows = int(sys.argv[2]) if len(sys.argv) >= 3 else 30
    wb = load_workbook(str(path), read_only=True, data_only=True)
    print(f"\n=== {path.name} ===")
    print(f"Sheets: {wb.sheetnames}\n")
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n--- Sheet: {name} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                print(f"... (+{ws.max_row - max_rows} more rows)")
                break
            shown = [str(c) if c is not None else "" for c in row]
            print(" | ".join(shown).rstrip())
    return 0


if __name__ == "__main__":
    sys.exit(main())
